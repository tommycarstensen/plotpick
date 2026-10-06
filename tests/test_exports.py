"""Tests for exports.py: the R, Excel and LaTeX exports.

The R script is run with Rscript and the LaTeX compiled with pdflatex when
they are installed; the other checks run everywhere.  CI's exports job
installs both and sets PLOTPICK_EXPORT_TOOLS, so that a missing tool fails
there instead of skipping.
"""

import io
import os
import shutil
import subprocess

import openpyxl
import pandas as pd
import pytest

from exports import (
    SHEET_MAX,
    dataframe_to_excel,
    dataframe_to_latex,
    dataframe_to_r,
    sheet_names,
)

ZIP_LABEL = "batch.zip/trial_2019.pdf p.3 Figure 2"
REQUIRED = bool(os.environ.get("PLOTPICK_EXPORT_TOOLS"))


def missing(tool: str) -> bool:
    """Whether to skip a check that needs `tool`: never when CI requires it."""
    return shutil.which(tool) is None and not REQUIRED


def table() -> pd.DataFrame:
    """Rows as the app builds them: text, numbers, a gap and awkward text."""
    return pd.DataFrame([
        {"source": ZIP_LABEL, "group": 'He said "hi"', "timepoint": "Baseline",
         "mean": 1.5, "error_type": "SD"},
        {"source": ZIP_LABEL, "group": "C:\\path & 50% #1", "timepoint": 6,
         "mean": 2.0, "error_type": "SE_mean"},
        {"source": "paper.pdf p.1 Table 1", "group": "\u00b5g/L", "timepoint": None,
         "mean": None, "error_type": None},
    ])


def test_r_script_has_a_comma_after_every_column():
    lines = dataframe_to_r(table()).splitlines()
    columns = [line for line in lines if line.startswith("  `")]
    assert len(columns) == 5
    assert all(line.endswith("),") for line in columns)


def test_r_script_quotes_text_and_leaves_numbers_bare():
    script = dataframe_to_r(table())
    group = '`group` = c("He said \\"hi\\"", "C:\\\\path & 50% #1", "\u00b5g/L")'
    assert group in script
    assert '`timepoint` = c("Baseline", "6", NA)' in script
    assert "`mean` = c(1.5, 2.0, NA)" in script


@pytest.mark.skipif(missing("Rscript"), reason="R is not installed")
def test_r_script_runs_in_r(tmp_path):
    path = tmp_path / "plotpick.R"
    path.write_text(
        dataframe_to_r(table())
        + 'stopifnot(nrow(dat) == 3, is.numeric(dat$mean), is.na(dat$mean[3]),\n'
        + '          dat$group[1] == "He said \\"hi\\"",\n'
        + '          dat$timepoint[2] == "6", dat$error_type[2] == "SE_mean")\n',
        encoding="utf-8",
    )
    run = subprocess.run(
        ["Rscript", str(path)], capture_output=True, text=True, check=False,
    )
    assert run.returncode == 0, run.stderr


def test_sheet_names_are_valid_and_unique():
    labels = [
        ZIP_LABEL,
        "a/b:c*d?e[f]g\\h'",
        "x" * 40 + " p.3 Figure 2",
        "y" * 40 + " p.3 Figure 2",
        "ALL",
        "",
    ]
    names = sheet_names(labels)
    assert list(names) == labels
    for name in names.values():
        assert 0 < len(name) <= SHEET_MAX
        assert not set(name) & set("[]:*?/\\'")
    lowered = [name.lower() for name in names.values()]
    assert len(set(lowered)) == len(lowered)
    assert "all" not in lowered
    assert names[ZIP_LABEL].endswith("p.3 Figure 2")


def test_excel_has_a_sheet_per_source_for_zip_labels():
    workbook = openpyxl.load_workbook(io.BytesIO(dataframe_to_excel(table())))
    assert workbook.sheetnames[0] == "All"
    assert len(workbook.sheetnames) == 3
    assert workbook["All"].max_row == 4


def test_latex_escapes_special_characters():
    latex = dataframe_to_latex(table())
    assert "error\\_type" in latex
    assert "SE\\_mean" in latex
    assert "50\\% \\#1" in latex
    assert "NaN" not in latex


@pytest.mark.skipif(missing("pdflatex"), reason="LaTeX is not installed")
def test_latex_compiles(tmp_path):
    (tmp_path / "table.tex").write_text(
        "\\documentclass{article}\n\\usepackage{lmodern}\n"
        "\\usepackage[T1]{fontenc}\n\\usepackage{textcomp}\n"
        "\\usepackage{booktabs}\n"
        "\\begin{document}\n" + dataframe_to_latex(table()) + "\\end{document}\n",
        encoding="utf-8",
    )
    run = subprocess.run(
        ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "table.tex"],
        cwd=tmp_path, capture_output=True, text=True, check=False,
    )
    assert run.returncode == 0, run.stdout[-2000:]
